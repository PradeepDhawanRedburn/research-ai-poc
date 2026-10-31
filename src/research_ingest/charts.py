"""Extract series data from native Word/Excel charts (DrawingML chart parts)."""
from __future__ import annotations

import datetime as dt
import re

from .package import NS, DocxPackage

C = NS["c"]
A = NS["a"]


def _c(tag):
    return "{%s}%s" % (C, tag)


def _text_of(el) -> str | None:
    if el is None:
        return None
    parts = [t.text or "" for t in el.iter("{%s}t" % A)]
    s = "".join(parts).strip()
    return s or None


def _excel_serial_to_date(v: float) -> str:
    base = dt.date(1899, 12, 30)
    try:
        return (base + dt.timedelta(days=int(v))).isoformat()
    except (OverflowError, ValueError):
        return str(v)


def _cache(ref_el):
    """Return (formula, format_code, [values]) from a numRef/strRef/numLit/strLit."""
    if ref_el is None:
        return None, None, []
    f = ref_el.find(".//" + _c("f"))  # <c:f> sits inside numRef/strRef, one level below the val/cat/tx element
    formula = f.text if f is not None else None
    cache = None
    for name in ("numCache", "strCache", "numLit", "strLit"):
        cache = ref_el.find(".//" + _c(name))
        if cache is not None:
            break
    if cache is None:
        return formula, None, []
    fmt_el = cache.find(_c("formatCode"))
    fmt = fmt_el.text if fmt_el is not None else None
    count_el = cache.find(_c("ptCount"))
    count = int(count_el.get("val")) if count_el is not None else None
    vals: dict[int, str] = {}
    for pt in cache.findall(_c("pt")):
        v = pt.find(_c("v"))
        vals[int(pt.get("idx"))] = v.text if v is not None else None
    n = count if count is not None else (max(vals) + 1 if vals else 0)
    out = [vals.get(i) for i in range(n)]
    return formula, fmt, out


def _is_date_format(fmt: str | None) -> bool:
    if not fmt:
        return False
    f = fmt.lower()
    return bool(re.search(r"[dmy]{2,}|mmm|yy", f)) and "0.0" not in f


def _to_number(v):
    if v is None:
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return v


RANGE_RE = re.compile(r"^\s*'?(?P<sheet>[^'!]+)'?!\$?(?P<c1>[A-Za-z]+)\$?(?P<r1>\d+)(?::\$?(?P<c2>[A-Za-z]+)\$?(?P<r2>\d+))?\s*$")


def _workbook(pkg: DocxPackage, xlsx_part: str):
    """The embedded workbook behind a chart, opened once per package (openpyxl, cached values only)."""
    cache = getattr(pkg, "_workbooks", None)
    if cache is None:
        cache = pkg._workbooks = {}
    if xlsx_part not in cache:
        try:
            import io

            from openpyxl import load_workbook

            cache[xlsx_part] = load_workbook(io.BytesIO(pkg.read(xlsx_part)), read_only=True, data_only=True)
        except Exception:  # noqa: BLE001 - a missing or unreadable workbook means no fallback, not a failure
            cache[xlsx_part] = None
    return cache[xlsx_part]


def workbook_values(pkg: DocxPackage, xlsx_part: str | None, formula: str | None) -> list:
    """Cells of a chart reference formula ("Sheet1!$B$2:$B$4") read from the embedded workbook, in row-major
    order; [] when there is no workbook, no formula, or the range cannot be resolved."""
    m = RANGE_RE.match(formula or "")
    if not xlsx_part or m is None:
        return []
    wb = _workbook(pkg, xlsx_part)
    if wb is None or m.group("sheet") not in wb.sheetnames:
        return []
    ws = wb[m.group("sheet")]
    from openpyxl.utils import column_index_from_string

    c1, r1 = column_index_from_string(m.group("c1").upper()), int(m.group("r1"))
    c2 = column_index_from_string(m.group("c2").upper()) if m.group("c2") else c1
    r2 = int(m.group("r2")) if m.group("r2") else r1
    out = []
    for row in ws.iter_rows(min_row=min(r1, r2), max_row=max(r1, r2), min_col=min(c1, c2), max_col=max(c1, c2),
                            values_only=True):
        for v in row:
            if isinstance(v, (dt.date, dt.datetime)):
                v = v.date().isoformat() if isinstance(v, dt.datetime) else v.isoformat()
            out.append(v)
    return out


def parse_chart(pkg: DocxPackage, part: str) -> dict:
    root = pkg.xml(part)
    chart = root.find(_c("chart"))
    out: dict = {"part": part, "title": None, "kinds": [], "series": [], "axes": {}, "external_data": None}
    ext = root.find(_c("externalData"))
    if ext is not None:
        rid = ext.get("{%s}id" % NS["r"])
        out["external_data"] = pkg.rels(part).target(rid) if rid else None
    if chart is None:
        return out
    out["title"] = _text_of(chart.find(_c("title")))
    plot = chart.find(_c("plotArea"))
    if plot is None:
        return out
    # axis titles
    for ax_tag in ("catAx", "valAx", "dateAx", "serAx"):
        for i, ax in enumerate(plot.findall(_c(ax_tag))):
            t = _text_of(ax.find(_c("title")))
            pos = ax.find(_c("axPos"))
            key = f"{ax_tag}{i}"
            out["axes"][key] = {"title": t, "position": pos.get("val") if pos is not None else None,
                                "id": (ax.find(_c("axId")).get("val") if ax.find(_c("axId")) is not None else None)}
    for child in plot:
        tag = child.tag.split("}")[1]
        if not tag.endswith("Chart"):
            continue
        out["kinds"].append(tag)
        grouping = child.find(_c("grouping"))
        bar_dir = child.find(_c("barDir"))
        for ser in child.findall(_c("ser")):
            name = None
            tx = ser.find(_c("tx"))
            if tx is not None:
                _, _, tv = _cache(tx.find(_c("strRef")))
                direct = tx.find(_c("v"))
                name = tv[0] if tv else (direct.text if direct is not None else _text_of(tx))
            cat_f, cat_fmt, cats = _cache(ser.find(_c("cat")))
            val_f, val_fmt, vals = _cache(ser.find(_c("val")))
            # scatter charts use xVal / yVal (only when they exist, or an uncached <c:val> loses its formula)
            if not vals and ser.find(_c("yVal")) is not None:
                cat_f, cat_fmt, cats = _cache(ser.find(_c("xVal")))
                val_f, val_fmt, vals = _cache(ser.find(_c("yVal")))
            # no cached values (Word wrote the chart without a cache, or it was stripped): the embedded workbook
            # still has them, addressed by the same reference formulas
            source = "cache"
            if not [v for v in vals if v not in (None, "")] and val_f and out["external_data"]:
                wb_vals = workbook_values(pkg, out["external_data"], val_f)
                if wb_vals:
                    vals, source = [str(v) if v is not None else None for v in wb_vals], "workbook"
                    if not [c for c in cats if c not in (None, "")]:
                        cats = [("" if c is None else str(c)) for c in workbook_values(pkg, out["external_data"], cat_f)]
                    if not name and tx is not None:
                        tx_f = tx.find(_c("strRef") + "/" + _c("f"))
                        wb_name = workbook_values(pkg, out["external_data"], tx_f.text if tx_f is not None else None)
                        name = str(wb_name[0]) if wb_name and wb_name[0] is not None else name
            if _is_date_format(cat_fmt):
                cats = [(_excel_serial_to_date(float(c)) if c not in (None, "") else None) for c in cats]
            numeric_vals = [_to_number(v) for v in vals]
            out["series"].append({
                "name": name,
                "values_source": source,
                "chart_kind": tag,
                "grouping": grouping.get("val") if grouping is not None else None,
                "bar_direction": bar_dir.get("val") if bar_dir is not None else None,
                "categories_formula": cat_f,
                "values_formula": val_f,
                "category_format": cat_fmt,
                "value_format": val_fmt,
                "categories": cats,
                "values": numeric_vals,
                "n_points": len(numeric_vals),
            })
    return out


def chart_summary_text(chart: dict) -> str:
    """Compact, lossless-enough text for an LLM: per series, range and endpoints."""
    lines = []
    kinds = ", ".join(sorted(set(chart["kinds"]))) or "chart"
    if chart.get("title"):
        lines.append(f"Chart title: {chart['title']}")
    lines.append(f"Chart type: {kinds}; {len(chart['series'])} series")
    for ax in chart["axes"].values():
        if ax.get("title"):
            lines.append(f"Axis ({ax['position']}): {ax['title']}")
    for s in chart["series"]:
        nums = [(c, v) for c, v in zip(s["categories"], s["values"]) if isinstance(v, (int, float))]
        if not nums:
            lines.append(f"- series {s['name'] or '(unnamed)'}: no cached values")
            continue
        vals = [v for _, v in nums]
        first_c, first_v = nums[0]
        last_c, last_v = nums[-1]
        mn = min(nums, key=lambda cv: cv[1])
        mx = max(nums, key=lambda cv: cv[1])
        if len(nums) <= 30:
            pairs = ", ".join(f"{c}: {v:g}" for c, v in nums)
            lines.append(f"- series {s['name'] or '(unnamed)'} ({len(nums)} points): {pairs}")
            continue
        tail = ", ".join(f"{c}: {v:g}" for c, v in nums[-6:])
        lines.append(
            f"- series {s['name'] or '(unnamed)'} ({len(nums)} points, {first_c} to {last_c}): "
            f"first {first_v:g} at {first_c}; min {mn[1]:g} at {mn[0]}; max {mx[1]:g} at {mx[0]}; "
            f"last six: {tail}"
        )
    return "\n".join(lines)
